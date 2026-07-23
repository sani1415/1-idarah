# -*- coding: utf-8 -*-
"""দপ্তর হিসাব — স্ট্যান্ডার্ড .xlsx টেমপ্লেট (অ্যাপ ফাইল-আপলোড ইমপোর্ট)।

Source of truth for in-app «টেমপ্লেট ডাউনলোড».
Writes:
  app/madrasa/templates/daftar-hisab-template-v2.xlsx  (served by the app)
  output/daftar-hisab-template.xlsx                 (local copy)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from pathlib import Path
from datetime import date, timedelta

ROOT = Path(__file__).resolve().parents[1]
OUTS = [
    ROOT / "app" / "madrasa" / "templates" / "daftar-hisab-template-v2.xlsx",
    ROOT / "output" / "daftar-hisab-template.xlsx",
]

header_fill = PatternFill("solid", fgColor="1A1208")
header_font = Font(name="Calibri", bold=True, color="F5E6C8", size=11)
title_font = Font(name="Calibri", bold=True, size=14, color="1A1208")
note_font = Font(name="Calibri", size=10, color="5C5346")
sample_fill = PatternFill("solid", fgColor="FFFBF2")
dash_fill = PatternFill("solid", fgColor="F7F1E6")
auto_fill = PatternFill("solid", fgColor="F0F7F4")
thin = Border(
    left=Side(style="thin", color="D9D0C3"),
    right=Side(style="thin", color="D9D0C3"),
    top=Side(style="thin", color="D9D0C3"),
    bottom=Side(style="thin", color="D9D0C3"),
)
wrap = Alignment(wrap_text=True, vertical="top")

# ব্যয়: E=পরিমাণ, G=একক মূল্য, H=মোট
TOTAL_FORMULA = '=IF(AND(E{r}<>"",G{r}<>""),E{r}*G{r},"")'
FORMULA_ROWS = 500
DATA_ROWS = 5000
HIJRI_DATE_FORMAT = "[$-1170401]yyyy-mm-dd"
WORKBOOK_ID = "daftar-hisab-v3-20260723"

# Excel-এর native Umm al-Qura calendar-এর সঙ্গে মিল রেখে নমুনা হিজরী তারিখকে
# আসল date serial-এ রূপান্তর করা হয়। প্রদর্শন ও app import—দুটিই এই calendar ব্যবহার করে।
HIJRI_MONTH_STARTS = {
    (1447, 8): date(2026, 1, 20),
    (1447, 9): date(2026, 2, 18),
    (1447, 10): date(2026, 3, 20),
    (1448, 1): date(2026, 6, 16),
    (1448, 2): date(2026, 7, 15),
}


def excel_date(value):
    """Template sample input → real Gregorian-backed Excel date."""
    if isinstance(value, date):
        return value
    parts = [int(x) for x in str(value).split("-")]
    if parts[0] >= 1700:
        return date(parts[0], parts[1], parts[2])
    start = HIJRI_MONTH_STARTS[(parts[0], parts[1])]
    return start + timedelta(days=parts[2] - 1)


def format_date_column(ws, column, end_row=DATA_ROWS):
    for row_idx in range(2, end_row + 1):
        ws.cell(row_idx, column).number_format = HIJRI_DATE_FORMAT


def add_hijri_helpers(ws):
    """Hidden helper columns make Hijri month/category SUMIFS auditable."""
    ws["L1"] = "হিজরী বছর (অটো)"
    ws["M1"] = "হিজরী মাস নং (অটো)"
    for cell in (ws["L1"], ws["M1"]):
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for row_idx in range(2, DATA_ROWS + 1):
        ws.cell(row_idx, 12, f'=IF(B{row_idx}="","",IFERROR(VALUE(TEXT(B{row_idx},"[$-1170401]yyyy")),""))')
        ws.cell(row_idx, 13, f'=IF(B{row_idx}="","",IFERROR(VALUE(TEXT(B{row_idx},"[$-1170401]mm")),""))')
    ws.column_dimensions["L"].hidden = True
    ws.column_dimensions["M"].hidden = True


def style_header(ws, cols, widths):
    for i, (name, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(1, i, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"


def style_data_row(ws, row_idx, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row_idx, c)
        cell.border = thin
        cell.fill = sample_fill
        cell.alignment = Alignment(vertical="center")


def set_total_auto(ws, row_idx):
    cell = ws.cell(row_idx, 8)
    cell.value = TOTAL_FORMULA.format(r=row_idx)
    cell.number_format = '#,##0.##'
    cell.fill = auto_fill
    cell.border = thin
    cell.alignment = Alignment(vertical="center")


def set_total_manual(ws, row_idx, amount):
    cell = ws.cell(row_idx, 8, amount)
    cell.number_format = '#,##0.##'
    cell.fill = sample_fill
    cell.border = thin
    cell.alignment = Alignment(vertical="center")


def main():
    wb = Workbook()

    # ── নির্দেশনা ──
    ws = wb.active
    ws.title = "নির্দেশনা"
    ws["A1"] = "দপ্তর হিসাব — স্ট্যান্ডার্ড Excel টেমপ্লেট"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")
    lines = [
        ("এই ফাইল কী", "দৈনন্দিন সব লেনদেন এখানে রাখুন। দিনশেষে অ্যাপে আপলোড করলে অ্যাপ লেটেস্ট দেখাবে।"),
        ("শীট", "ব্যয় · আয় · করজ · পরিশোধ · বকেয়া · ড্যাশবোর্ড। ব্যয়/আয়/করজ/পরিশোধ অ্যাপে যায়; বকেয়া ও ড্যাশবোর্ড হিসাব দেখায়।"),
        ("ব্যয়", "মাতবাখ/মাদরাসা/তামিরাত/সাধারণ। বাকি কেনাকাটায় পরিশোধ=বাকি + সরবরাহকারী নাম।"),
        ("মোট টাকা", "পরিমাণ × একক মূল্য হলে মোট অটো। বিল হলে শুধু মোট লিখুন।"),
        ("বকেয়া", "এই শীটের «সরবরাহকারী» কলামে নাম লিখুন — ব্যয় ও পরিশোধ শীটের ড্রপডাউনে অটো আসবে। অবশিষ্ট = বাকি কেনা − পরিশোধ।"),
        ("পরিশোধ", "বকেয়া মেটানোর এন্ট্রি; সংশ্লিষ্ট হিসাব বিভাগ ও সরবরাহকারী বেছে নিন।"),
        ("আয় / করজ", "নিয়মিত আয় আলাদা; করজে দেওয়া/আদায় করজ শীটে।"),
        ("তারিখ", "ঘরটি আসল Excel Date; দেখা যাবে হিজরী সন-মাস-দিন। Date Filter/Sort কাজ করবে। নতুন তারিখ Excel date picker বা স্বীকৃত date entry দিয়ে দিন।"),
        ("অ্যাপে", "হিসাব → এক্সেল ↧ → এই ফাইল আপলোড। একই ফাইল আবার দিলে একই সারি আপডেট হবে; ডুপ্লিকেট হবে না।"),
    ]
    r = 3
    for a, b in lines:
        ws.cell(r, 1, a).font = Font(bold=True, size=11)
        ws.cell(r, 2, b).alignment = wrap
        ws.row_dimensions[r].height = 40
        r += 1
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 92
    ws["Z1"] = WORKBOOK_ID
    ws.column_dimensions["Z"].hidden = True

    # ── ব্যয় ──
    ws = wb.create_sheet("ব্যয়")
    exp_cols = [
        "হিসাব বিভাগ *", "তারিখ *", "খাত", "বিবরণ", "পরিমাণ", "মাপ",
        "একক মূল্য", "মোট টাকা (অটো)", "সরবরাহকারী", "রশিদ নং", "পরিশোধ",
    ]
    style_header(ws, exp_cols, [14, 14, 14, 26, 10, 10, 12, 16, 18, 12, 10])

    # qty×price rows → formula; bill-only → manual total
    qty_rows = [
        (2, ["মাতবাখ", excel_date("1447-10-02"), "কাঁচামাল", "চাল", 50, "কেজি", 70], "রহিম স্টোর", "১০১", "নগদ"),
        (3, ["মাতবাখ", excel_date("1447-10-03"), "কাঁচামাল", "ডাল", 20, "কেজি", 120], "রহিম স্টোর", "", "নগদ"),
        (4, ["মাদরাসা", excel_date("1447-10-04"), "স্টেশনারি", "খাতা ও কলম", 10, "প্যাকেট", 80], "বইঘর", "", "নগদ"),
        (5, ["তামিরাত", excel_date("1447-09-20"), "মেরামত", "ছাদের প্লাস্টার", 1, "পিস", 15000], "রাজমিস্ত্রি করিম", "", "বাকি"),
    ]
    for row_idx, left, supplier, receipt, pay in qty_rows:
        for c, val in enumerate(left, 1):
            ws.cell(row_idx, c, val)
        ws.cell(row_idx, 9, supplier)
        ws.cell(row_idx, 10, receipt)
        ws.cell(row_idx, 11, pay)
        style_data_row(ws, row_idx, 11)
        set_total_auto(ws, row_idx)

    bill_rows = [
        (6, ["সাধারণ", excel_date("1447-10-06"), "পরিবহন", "বাজার আনা-নেওয়া", "", "", ""], 500, "", "", "নগদ"),
        (7, ["মাদরাসা", excel_date("2026-07-10"), "বিদ্যুৎ", "বিদ্যুৎ বিল", "", "", ""], 4500, "পল্লী বিদ্যুৎ", "বিল-৮৮", "নগদ"),
    ]
    for row_idx, left, amount, supplier, receipt, pay in bill_rows:
        for c, val in enumerate(left, 1):
            ws.cell(row_idx, c, val)
        ws.cell(row_idx, 9, supplier)
        ws.cell(row_idx, 10, receipt)
        ws.cell(row_idx, 11, pay)
        style_data_row(ws, row_idx, 11)
        set_total_manual(ws, row_idx, amount)

    # খালি সারিতেও অটো ফর্মুলা — নতুন এন্ট্রিতে মোট নিজে আসবে
    for row_idx in range(8, FORMULA_ROWS + 1):
        set_total_auto(ws, row_idx)

    add_hijri_helpers(ws)
    format_date_column(ws, 2)
    ws.auto_filter.ref = "A1:M1"

    ws["O1"] = (
        "মোট টাকা = পরিমাণ × একক মূল্য (অটো)। "
        "বিল হলে পরিমাণ/মূল্য খালি রেখে মোট ঘরে সরাসরি টাকা লিখুন।"
    )
    ws["O1"].font = note_font
    ws.column_dimensions["O"].width = 55

    dv_book = DataValidation(type="list", formula1='"মাতবাখ,মাদরাসা,তামিরাত,সাধারণ"', allow_blank=False)
    ws.add_data_validation(dv_book)
    dv_book.add("A2:A5000")
    dv_unit = DataValidation(
        type="list",
        formula1='"কেজি,গ্রাম,লিটার,মিলি,পিস,প্যাকেট,বস্তা,ডজন,ফুট,মিটার"',
        allow_blank=True,
    )
    ws.add_data_validation(dv_unit)
    dv_unit.add("F2:F5000")
    dv_pay = DataValidation(type="list", formula1='"নগদ,বাকি"', allow_blank=True)
    ws.add_data_validation(dv_pay)
    dv_pay.add("K2:K5000")

    # ── আয় ──
    ws = wb.create_sheet("আয়")
    style_header(ws, ["তারিখ *", "টাকা *", "উৎস / বিবরণ"], [16, 14, 40])
    for i, row in enumerate([
        [excel_date("1447-10-01"), 50000, "বিকাশ অনুদান"],
        [excel_date("1447-10-05"), 12000, "ওয়াযিফা আদায় — শাওয়াল"],
        [excel_date("2026-07-10"), 3000, "নগদ দান"],
    ], start=2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(i, c, val)
            cell.border = thin
            cell.fill = sample_fill
    format_date_column(ws, 1)

    # ── করজ ──
    ws = wb.create_sheet("করজ")
    style_header(ws, ["ধরন *", "তারিখ *", "টাকা *", "খাত *", "বিবরণ"], [12, 14, 12, 22, 28])
    for i, row in enumerate([
        ["দেওয়া", excel_date("1447-08-15"), 20000, "করিম ভাই", "করজে হাসানা"],
        ["দেওয়া", excel_date("1447-09-01"), 10000, "মাদরাসা ফান্ড", "জরুরি করজ"],
        ["আদায়", excel_date("1447-10-07"), 5000, "করিম ভাই", ""],
        ["আদায়", excel_date("1447-10-11"), 2000, "মাদরাসা ফান্ড", ""],
    ], start=2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(i, c, val)
            cell.border = thin
            cell.fill = sample_fill
    format_date_column(ws, 2)
    dv_kind = DataValidation(type="list", formula1='"দেওয়া,আদায়"', allow_blank=False)
    ws.add_data_validation(dv_kind)
    dv_kind.add("A2:A5000")

    # ── পরিশোধ (বকেয়া মেটানোর এন্ট্রি) ──
    ws = wb.create_sheet("পরিশোধ")
    style_header(ws, ["হিসাব বিভাগ *", "তারিখ *", "সরবরাহকারী *", "টাকা *", "রশিদ নং", "মন্তব্য"], [16, 16, 24, 14, 14, 32])
    for i, row in enumerate([
        ["তামিরাত", excel_date("1447-10-15"), "রাজমিস্ত্রি করিম", 5000, "", "আংশিক পরিশোধ"],
    ], start=2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(i, c, val)
            cell.border = thin
            cell.fill = sample_fill
    format_date_column(ws, 2)
    ws["H1"] = "বিভাগ ও সরবরাহকারী—দুটিই সংশ্লিষ্ট বকেয়া কেনার সঙ্গে মিলতে হবে।"
    ws["H1"].font = note_font
    ws.column_dimensions["H"].width = 52
    dv_pay_book = DataValidation(type="list", formula1='"মাতবাখ,মাদরাসা,তামিরাত,সাধারণ"', allow_blank=False)
    ws.add_data_validation(dv_pay_book)
    dv_pay_book.add("A2:A5000")

    # ── বকেয়া (লাইভ তালিকা — সহজ ফর্মুলা, UNIQUE/FILTER নেই) ──
    ws = wb.create_sheet("বকেয়া")
    ws["A1"] = "সরবরাহকারী বকেয়া (লাইভ)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:E1")
    ws["A2"] = (
        "এই টেবিলের «সরবরাহকারী» কলাম = মাস্টার তালিকা। "
        "এখানে নাম লিখলে ব্যয় ও পরিশোধ শীটের ড্রপডাউনে চলে আসবে। "
        "অবশিষ্ট = বাকি কেনা − পরিশোধ।"
    )
    ws["A2"].font = note_font
    ws["A2"].alignment = wrap
    ws.merge_cells("A2:E3")
    ws.row_dimensions[2].height = 44

    summary = [
        (5, "মোট বাকি কেনা", '=SUMIF(ব্যয়!K:K,"বাকি",ব্যয়!H:H)'),
        (6, "মোট পরিশোধ", "=SUM(পরিশোধ!D:D)"),
        (7, "এখন মোট বাকি", "=B5-B6"),
    ]
    for row, label, formula in summary:
        a = ws.cell(row, 1, label)
        b = ws.cell(row, 2, formula)
        a.font = Font(bold=True, size=11)
        a.fill = dash_fill
        a.border = thin
        b.number_format = '#,##0'
        b.border = thin
        b.fill = auto_fill
        if row == 7:
            b.font = Font(bold=True, size=14, color="C1440E")

    headers = ["সরবরাহকারী", "বাকি কেনা", "পরিশোধ", "অবশিষ্ট বাকি", "বার"]
    for col, title in enumerate(headers, 1):
        cell = ws.cell(9, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin

    # মাস্টার তালিকা (ড্রপডাউনের সোর্স) + লাইভ হিসাব
    supplier_seed = ["রাজমিস্ত্রি করিম", "রহিম স্টোর", "বইঘর", "পল্লী বিদ্যুৎ"]
    suppliers = supplier_seed + [""] * (100 - len(supplier_seed))
    for i, name in enumerate(suppliers):
        r = 10 + i
        ws.cell(r, 1, name).border = thin
        ws.cell(r, 1).fill = sample_fill
        ws.cell(r, 2).value = (
            f'=IF(A{r}="","",SUMIFS(ব্যয়!$H$2:$H$500,ব্যয়!$I$2:$I$500,A{r},ব্যয়!$K$2:$K$500,"বাকি"))'
        )
        ws.cell(r, 3).value = (
            f'=IF(A{r}="","",SUMIF(পরিশোধ!$C$2:$C$500,A{r},পরিশোধ!$D$2:$D$500))'
        )
        ws.cell(r, 4).value = f'=IF(A{r}="","",B{r}-C{r})'
        ws.cell(r, 5).value = (
            f'=IF(A{r}="","",COUNTIFS(ব্যয়!$I$2:$I$500,A{r},ব্যয়!$K$2:$K$500,"বাকি"))'
        )
        for c in range(2, 6):
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = auto_fill
            if c < 5:
                ws.cell(r, c).number_format = '#,##0'

    ws["A111"] = (
        "টিপ: নতুন সরবরাহকারী → এখানে নাম লিখুন → ব্যয়/পরিশোধে ড্রপডাউন থেকে বেছে নিন।"
    )
    ws["A111"].font = note_font
    ws.merge_cells("A111:E111")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 8

    # ড্রপডাউন: বকেয়া!A10:A109 → ব্যয় (সরবরাহকারী) ও পরিশোধ
    try:
        wb.defined_names.delete("SupplierList")
    except (KeyError, AttributeError, TypeError):
        pass
    wb.defined_names.add(DefinedName(name="SupplierList", attr_text="'বকেয়া'!$A$10:$A$109"))

    dv_sup_exp = DataValidation(type="list", formula1="=SupplierList", allow_blank=True)
    dv_sup_exp.error = "বকেয়া শীটের তালিকা থেকে বেছে নিন, বা আগে সেখানে নাম যোগ করুন"
    dv_sup_exp.errorTitle = "সরবরাহকারী"
    dv_sup_exp.prompt = "বকেয়া শীটের তালিকা"
    dv_sup_exp.promptTitle = "সরবরাহকারী"
    wb["ব্যয়"].add_data_validation(dv_sup_exp)
    dv_sup_exp.add("I2:I5000")

    dv_sup_pay = DataValidation(type="list", formula1="=SupplierList", allow_blank=True)
    wb["পরিশোধ"].add_data_validation(dv_sup_pay)
    dv_sup_pay.add("C2:C5000")

    # ── ড্যাশবোর্ড ──
    ws = wb.create_sheet("ড্যাশবোর্ড")
    ws["A1"] = "সামারি (লাইভ)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:C1")

    labels = [
        (3, "মোট ব্যয় (সব বিভাগ)", "=SUM(ব্যয়!H:H)"),
        (4, "মোট আয়", "=SUM(আয়!B:B)"),
        (5, "করজ দেওয়া", '=SUMIF(করজ!A:A,"দেওয়া",করজ!C:C)'),
        (6, "করজ আদায়", '=SUMIF(করজ!A:A,"আদায়",করজ!C:C)'),
        (7, "করজ বাকি", "=B5-B6"),
        (8, "সরবরাহকারী বাকি", "=বকেয়া!B7"),
        (9, "নগদ (আনুমানিক)", "=B4-B3-B5+B6"),
    ]
    ws["A2"] = "বিবরণ"
    ws["B2"] = "টাকা"
    for c in (ws["A2"], ws["B2"]):
        c.fill = header_fill
        c.font = header_font
        c.border = thin
    for row, label, formula in labels:
        a = ws.cell(row, 1, label)
        b = ws.cell(row, 2, formula)
        a.border = thin
        b.border = thin
        a.fill = dash_fill
        b.number_format = '#,##0'
    ws["A11"] = "বিভাগ অনুযায়ী ব্যয়"
    ws["A11"].font = Font(bold=True, size=12)
    books = ["মাতবাখ", "মাদরাসা", "তামিরাত", "সাধারণ"]
    ws["A12"] = "হিসাব বিভাগ"
    ws["B12"] = "মোট"
    for c in (ws["A12"], ws["B12"]):
        c.fill = header_fill
        c.font = header_font
        c.border = thin
    for i, book in enumerate(books):
        rr = 13 + i
        ws.cell(rr, 1, book).border = thin
        cell = ws.cell(rr, 2, f"=SUMIF(ব্যয়!A:A,A{rr},ব্যয়!H:H)")
        cell.border = thin
        cell.number_format = "#,##0"
    ws["A18"] = (
        "বকেয়া তালিকায় নাম লিখুন → ব্যয়/পরিশোধে ড্রপডাউন। "
        "পরিশোধ «পরিশোধ» শীটে।"
    )
    ws["A18"].font = note_font
    ws.merge_cells("A18:C19")

    # ── নির্বাচিত হিজরী মাসে খাতভিত্তিক ব্যয় ──
    ws["A21"] = "হিজরী বছর"
    ws["B21"] = 1447
    ws["C21"] = "হিজরী মাস"
    ws["D21"] = "শাওয়াল"
    ws["J21"] = '=MATCH(D21,K2:K13,0)'
    for c in (ws["A21"], ws["C21"]):
        c.fill = header_fill
        c.font = header_font
        c.border = thin
    for c in (ws["B21"], ws["D21"]):
        c.fill = sample_fill
        c.border = thin
        c.font = Font(bold=True, color="1A1208")

    dashboard_years = list(range(1447, 1451))
    dv_year = DataValidation(type="list", formula1='"' + ",".join(str(y) for y in dashboard_years) + '"', allow_blank=False)
    ws.add_data_validation(dv_year)
    dv_year.add("B21")
    hijri_months = [
        "মুহাররম", "সফর", "রবিউল আউয়াল", "রবিউস সানি", "জুমাদাল উলা", "জুমাদাল উখরা",
        "রজব", "শাবান", "রমজান", "শাওয়াল", "জিলকদ", "জিলহজ",
    ]
    dv_month = DataValidation(type="list", formula1='"' + ",".join(hijri_months) + '"', allow_blank=False)
    ws.add_data_validation(dv_month)
    dv_month.add("D21")
    for i, month_name in enumerate(hijri_months, start=2):
        ws.cell(i, 11, month_name)
        ws.cell(i, 12, i - 1)
    ws.column_dimensions["J"].hidden = True
    ws.column_dimensions["K"].hidden = True
    ws.column_dimensions["L"].hidden = True

    ws["A23"] = "নির্বাচিত মাসে খাতভিত্তিক ব্যয়"
    ws["A23"].font = Font(bold=True, size=12)
    ws.merge_cells("A23:B23")
    ws["A24"] = "খাত"
    ws["B24"] = "মোট টাকা"
    for c in (ws["A24"], ws["B24"]):
        c.fill = header_fill
        c.font = header_font
        c.border = thin
    categories = [
        "বড় বাজার", "কাঁচা বাজার", "দস্তরখান ভবন", "রান্নাঘর সরবরাহ",
        "বিদ্যুৎ", "গ্যাস/জ্বালানি", "শিক্ষা উপকরণ", "বেতন", "রক্ষণাবেক্ষণ",
        "পরিবহন", "চিকিৎসা", "অন্যান্য", "কাঁচামাল", "স্টেশনারি", "মেরামত",
    ]
    for i, category in enumerate(categories, start=25):
        ws.cell(i, 1, category)
        ws.cell(i, 1).border = thin
        ws.cell(i, 1).fill = sample_fill
        amount = ws.cell(
            i, 2,
            f'=SUMIFS(\'ব্যয়\'!$H$2:$H${DATA_ROWS},\'ব্যয়\'!$L$2:$L${DATA_ROWS},$B$21,'
            f'\'ব্যয়\'!$M$2:$M${DATA_ROWS},$J$21,\'ব্যয়\'!$C$2:$C${DATA_ROWS},A{i})',
        )
        amount.border = thin
        amount.fill = auto_fill
        amount.number_format = "#,##0"
    total_row = 25 + len(categories)
    ws.cell(total_row, 1, "মোট")
    ws.cell(total_row, 2, f"=SUM(B25:B{total_row - 1})")
    for c in (ws.cell(total_row, 1), ws.cell(total_row, 2)):
        c.fill = dash_fill
        c.border = thin
        c.font = Font(bold=True, color="1A1208")
    ws.cell(total_row, 2).number_format = "#,##0"
    ws.cell(total_row + 2, 1, "নতুন/নিজস্ব খাত হলে উপরের খাতের নাম বদলে লিখুন; হিসাব অটো আপডেট হবে।")
    ws.cell(total_row + 2, 1).font = note_font
    ws.merge_cells(start_row=total_row + 2, start_column=1, end_row=total_row + 3, end_column=3)
    ws.cell(total_row + 2, 1).alignment = wrap

    # ── নির্বাচিত হিজরী বছরের ১২ মাস × হিসাব বিভাগ (উপরের পাশের compact block) ──
    ws["D1"] = '="হিজরী "&B21&" বর্ষের মাসভিত্তিক বিভাগীয় ব্যয়"'
    ws["D1"].font = Font(bold=True, size=12)
    ws.merge_cells("D1:I1")
    month_headers = ["মাস", "মাতবাখ", "মাদরাসা", "তামিরাত", "সাধারণ", "মাসের মোট"]
    for col, title in enumerate(month_headers, start=4):
        cell = ws.cell(2, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for month_no, month_name in enumerate(hijri_months, start=1):
        rr = month_no + 2
        ws.cell(rr, 4, month_name)
        ws.cell(rr, 4).border = thin
        ws.cell(rr, 4).fill = sample_fill
        for col, book in enumerate(books, start=5):
            cell = ws.cell(
                rr, col,
                f'=SUMIFS(\'ব্যয়\'!$H$2:$H${DATA_ROWS},\'ব্যয়\'!$L$2:$L${DATA_ROWS},$B$21,'
                f'\'ব্যয়\'!$M$2:$M${DATA_ROWS},{month_no},'
                f'\'ব্যয়\'!$A$2:$A${DATA_ROWS},"{book}")',
            )
            cell.border = thin
            cell.fill = auto_fill
            cell.number_format = "#,##0"
        ws.cell(rr, 9, f"=SUM(E{rr}:H{rr})")
        ws.cell(rr, 9).border = thin
        ws.cell(rr, 9).fill = dash_fill
        ws.cell(rr, 9).number_format = "#,##0"

    grand_total_row = 15
    ws.cell(grand_total_row, 4, "বর্ষের মোট")
    for col in range(5, 10):
        letter = get_column_letter(col)
        ws.cell(grand_total_row, col, f"=SUM({letter}3:{letter}14)")
        ws.cell(grand_total_row, col).number_format = "#,##0"
    for col in range(4, 10):
        cell = ws.cell(grand_total_row, col)
        cell.fill = dash_fill
        cell.border = thin
        cell.font = Font(bold=True, color="1A1208")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 20
    for col in ("E", "F", "G", "H", "I"):
        ws.column_dimensions[col].width = 14

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    for out in OUTS:
        out.parent.mkdir(parents=True, exist_ok=True)
        try:
            wb.save(out)
            print("saved:", out.as_posix())
        except PermissionError:
            alt = out.with_name(out.stem + "-latest.xlsx")
            wb.save(alt)
            print("locked, saved alt:", alt.as_posix())


if __name__ == "__main__":
    main()
