# -*- coding: utf-8 -*-
"""Upgrade a populated daftar workbook to total-input/unit-price-auto format."""
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\sanim\Downloads\daftar-hisab-template (10).xlsx")
LATEST = ROOT / "app" / "madrasa" / "templates" / "daftar-hisab-template-v2.xlsx"
OUTPUT = ROOT / "output" / "daftar-hisab-template-updated-with-data.xlsx"
UNIT_PRICE_FORMULA = '=IF(AND(E{r}<>"",E{r}<>0,H{r}<>""),H{r}/E{r},"")'


def copy_cell_style(source, target):
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def recover_amount(formula_cell, cached_cell, quantity, old_unit_price):
    raw = formula_cell.value
    if isinstance(raw, (int, float)):
        return raw
    if isinstance(cached_cell.value, (int, float)):
        return cached_cell.value
    if isinstance(quantity, (int, float)) and isinstance(old_unit_price, (int, float)):
        return quantity * old_unit_price
    return None


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not LATEST.exists():
        raise FileNotFoundError(LATEST)

    workbook = load_workbook(SOURCE, data_only=False)
    cached = load_workbook(SOURCE, data_only=True)
    latest = load_workbook(LATEST, data_only=False)
    expense = workbook.worksheets[1]
    cached_expense = cached.worksheets[1]
    latest_expense = latest.worksheets[1]

    before_total = 0.0
    recovered_count = 0
    unresolved_rows = []
    for row in range(2, 5001):
        quantity = expense.cell(row, 5).value
        old_unit_price = expense.cell(row, 7).value
        amount = recover_amount(
            expense.cell(row, 8), cached_expense.cell(row, 8), quantity, old_unit_price
        )
        meaningful = any(
            expense.cell(row, col).value not in (None, "")
            for col in (2, 4, 5, 6, 7, 9, 10, 11)
        )
        if amount is not None:
            before_total += float(amount)
            recovered_count += 1
        elif meaningful and (expense.cell(row, 2).value or expense.cell(row, 4).value):
            unresolved_rows.append(row)

        expense.cell(row, 8).value = amount
        expense.cell(row, 7).value = UNIT_PRICE_FORMULA.format(r=row)
        copy_cell_style(latest_expense.cell(row, 7), expense.cell(row, 7))
        copy_cell_style(latest_expense.cell(row, 8), expense.cell(row, 8))

    expense["G1"].value = latest_expense["G1"].value
    expense["H1"].value = latest_expense["H1"].value
    copy_cell_style(latest_expense["G1"], expense["G1"])
    copy_cell_style(latest_expense["H1"], expense["H1"])
    expense["O1"].value = latest_expense["O1"].value
    copy_cell_style(latest_expense["O1"], expense["O1"])
    expense.freeze_panes = latest_expense.freeze_panes
    expense.column_dimensions["L"].hidden = latest_expense.column_dimensions["L"].hidden
    expense.column_dimensions["M"].hidden = latest_expense.column_dimensions["M"].hidden

    instructions = workbook.worksheets[0]
    latest_instructions = latest.worksheets[0]
    for row in range(1, latest_instructions.max_row + 1):
        if latest_instructions.cell(row, 1).value == "মোট টাকা":
            for col in (1, 2):
                instructions.cell(row, col).value = latest_instructions.cell(row, col).value
                copy_cell_style(latest_instructions.cell(row, col), instructions.cell(row, col))
            break

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    workbook.close()
    cached.close()
    latest.close()
    print(f"saved={OUTPUT}")
    print(f"recovered_count={recovered_count}")
    print(f"recovered_total={before_total:.6f}")
    print("unresolved_rows=" + ",".join(str(x) for x in unresolved_rows))


if __name__ == "__main__":
    main()
